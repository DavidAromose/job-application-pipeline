"""Supporting statement / cover letter for a role, built from the CV's jd_match.

Why: NHS, councils, schools and many large employers shortlist by scoring the
"supporting information" box against the person specification, criterion by
criterion - a CV alone rarely gets past that. This writes a statement that walks
through the criteria in the advert's order, giving the TRUE evidence from the
master resume for each one the candidate meets.

Honesty rules (same as the CV):
  * Evidence comes only from jd_match / the master - nothing invented.
  * A criterion with no evidence is never claimed. It is either left out, or (for an
    essential one) acknowledged plainly with a genuine willingness to learn/train -
    never with an invented example.
  * No personality puffery ("passionate", "dynamic"...), no undergraduate-degree pitch.
"""
from __future__ import annotations

import json

import anthropic
from docx import Document
from docx.shared import Inches, Pt

from config import CONFIG
from pipeline.scraper import JobPosting

STATEMENT_SYSTEM = """You write the supporting statement (or short cover letter) for a UK job \
application. You are given the job advert, the candidate's MASTER RESUME (the only source of \
truth) and a requirement-by-requirement evidence map.

Rules:
- 300-450 words. Plain, specific, first person, British English.
- Open with one sentence saying which role this is for and one concrete, TRUE reason the \
candidate fits it.
- Then go through the advert's criteria in the advert's own order, using its key words as \
short signposts where helpful, and give the TRUE evidence for each criterion that has some - \
name the role/employer where it happened.
- Never claim anything the master does not state. Do not invent examples, numbers, systems, \
duties or qualifications. Do not add personality words (passionate, dynamic, calm, driven...).
- For an ESSENTIAL criterion with no evidence, write one honest sentence that addresses it \
without pretending (for example what related thing the candidate has done, or that they \
would complete the employer's training) - only if that related thing is true. Skip \
desirable gaps silently.
- Do not mention the undergraduate degree or its grade. Mention the MRes at most once and \
only if relevant.
- Finish with availability (only if stated in the master/candidate details) and a one-line close.
- Never use em dashes or en dashes (or a hyphen used as a dash). Write ranges as "February to June 2026".
- Return ONLY the statement text, no headings, no placeholders, no sign-off name."""


def write_statement(job: JobPosting, resume, master_text: str) -> str:
    client = anthropic.Anthropic(api_key=CONFIG.anthropic_api_key)
    app = CONFIG.applicant
    candidate = {
        "availability": app.availability or "",
        "right_to_work": app.right_to_work or "",
    }
    prompt = f"""JOB: {job.title} at {job.company}

ADVERT:
{job.description[:9000]}

EVIDENCE MAP (from the CV step):
{json.dumps(resume.jd_match, ensure_ascii=False, indent=1)}

CANDIDATE DETAILS (only use if non-empty): {json.dumps(candidate)}

MASTER RESUME (source of truth):
{master_text[:22000]}"""
    msg = client.messages.create(
        model=CONFIG.anthropic_model,
        max_tokens=3000,
        system=STATEMENT_SYSTEM,
        messages=[{"role": "user", "content": prompt}],
    )
    return next((b.text for b in msg.content if b.type == "text"), "").strip()


def build_statement_docx(text: str, job: JobPosting, out_path) -> str:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.9)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    app = CONFIG.applicant
    h = doc.add_paragraph()
    r = h.add_run(app.name)
    r.bold = True
    r.font.size = Pt(14)
    bits = [b for b in (app.location, app.email, app.phone) if b]
    if bits:
        doc.add_paragraph(" | ".join(bits)).runs[0].font.size = Pt(9)
    t = doc.add_paragraph()
    t.add_run(f"Supporting statement: {job.title}, {job.company}").bold = True
    from pipeline.resume import _no_dashes
    text = _no_dashes(text)
    for para in [p.strip() for p in text.split("\n") if p.strip()]:
        doc.add_paragraph(para)
    out_path = str(out_path)
    doc.save(out_path)
    return out_path
