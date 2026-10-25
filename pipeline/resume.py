"""Tailor the master resume to a job description via the Anthropic API.

Hard rules enforced here (both in the prompt and by code afterwards):
  * Never fabricate metrics, dates, employers, titles or achievements.
  * Only rephrase and reorder bullets/content that already exist in the master.
  * Prefer the job description's action verbs where it reads naturally.
  * Keep light-to-medium keyword overlap for ATS - not keyword stuffing.

The master resume is treated as a *source document*: the model selects a subset
of existing content, reorders it, and lightly rewords it into a one-to-two page
job-facing CV. A post-generation check flags any new numbers or any bullet that
does not map back to something in the master.
"""
from __future__ import annotations

import difflib
import json
import re
from dataclasses import dataclass, field

import anthropic
from docx import Document
from docx.shared import Pt

from config import CONFIG
from pipeline.scraper import JobPosting

NUM_RE = re.compile(r"\b\d[\d,.]*\s?(?:%|percent|k\b|m\b|bn\b|\+)?", re.IGNORECASE)
SIMILARITY_FLOOR = 0.55  # difflib ratio below which a bullet is "not from the master"


# --------------------------------------------------------------------------- #
# 1. Extract the master resume into plain structured blocks
# --------------------------------------------------------------------------- #
@dataclass
class Block:
    kind: str  # "heading" | "subheading" | "bullet" | "text"
    text: str


def _classify(paragraph) -> str:
    style = ((paragraph.style.name if paragraph.style else "") or "").lower()
    text = paragraph.text.strip()
    if not text:
        return "empty"
    if "heading 1" in style or "title" in style:
        return "heading"
    if "heading" in style:
        return "subheading"
    if style.startswith("list") or text[:2] in {"- ", "• ", "* "}:
        return "bullet"
    if re.match(r"^\d+\.\s+[A-Z]", text) and len(text) < 60:
        return "heading"
    runs_bold = [r for r in paragraph.runs if r.text.strip()]
    if runs_bold and all(r.bold for r in runs_bold) and len(text) < 120:
        return "subheading"
    return "text"


def extract_master(path=None) -> list[Block]:
    path = path or CONFIG.master_resume_path
    doc = Document(str(path))
    blocks: list[Block] = []
    for para in doc.paragraphs:
        kind = _classify(para)
        if kind == "empty":
            continue
        text = re.sub(r"^[-•*]\s+", "", para.text.strip())
        blocks.append(Block(kind=kind, text=text))
    return blocks


def master_corpus(blocks: list[Block]) -> list[str]:
    return [b.text for b in blocks if b.kind in {"bullet", "text", "subheading"}]


# --------------------------------------------------------------------------- #
# 2. Ask Claude to assemble a tailored resume (JSON)
# --------------------------------------------------------------------------- #
TAILOR_SYSTEM = """You are a meticulous resume editor. You tailor an existing \
master resume to a specific job description.

ABSOLUTE RULES:
- Never invent or alter facts: no new metrics, numbers, percentages, dates, \
employers, job titles, tools, or achievements. If the master does not contain a \
number, your output must not contain one.
- You may only rephrase and reorder content that already exists in the master.
- Every bullet you output must be traceable to a specific master bullet or \
sentence. Rewording for concision and impact is allowed; adding information is not.
- Where the job description uses a particular action verb for a responsibility \
the candidate genuinely has, prefer that verb - but only when it stays truthful.
- Aim for light-to-medium keyword overlap with the job description so it passes \
ATS keyword screens. Do not stuff keywords or list skills the master lacks.

RECRUITER-FIRST RULES (a recruiter spends ~10 seconds on a CV - make those seconds count):
1. ONE PAGE. At most 12 bullets in total, at most 3 per role, most job-relevant \
first. Omit roles that do not help for THIS job (list what you omitted in \
notes_for_human). Never misdate or hide a role's dates.
2. SUMMARY: 2-3 sentences. Open with the candidate's strongest TRUE strength for \
this exact job, using the job advert's own key words where truthful. Mention the \
MRes at most once, briefly, at the end - and for non-technical roles never lead \
with it. NEVER mention the undergraduate degree or its classification in the \
summary (it may appear only in the Education section).
3. MIRROR THE ADVERT: for each essential/desirable requirement in the job advert, \
if the master truthfully evidences it, word at least one bullet with the advert's \
own key nouns/verbs. If the master does NOT evidence it, do not claim it - record \
it as a GAP.
4. SKILLS: 6-10 items maximum, only ones the advert asks for or that are clearly \
relevant AND evidenced in the master. No filler like "adaptability" or "managing \
multiple commitments". For any role that is not technology/data/AI/research do NOT \
list Python, Git/GitHub, Jupyter, VS Code, TensorFlow, AWS or machine learning.
5. PROJECTS: only for technology/data/AI/research roles. For every other role output \
an empty projects list. Volunteering goes in "volunteering", NEVER under projects.
6. EDUCATION: MRes on one line (institution, dates, expected completion). List \
modules only for technology/data/AI/research roles. B.Agric on one line with its \
classification.
7. Never output placeholder text such as "[TBC" - skip any item that is unconfirmed.
7b. jd_match must list EVERY requirement in the advert's person specification / "what we're looking for" / "essential" and "desirable" lists (not just a sample). Set "essential": true for essential/required criteria and false for desirable ones; if the advert does not separate them, treat hard requirements ("must", "required", "experience of") as essential.
8. Do not repeat the same fact in the summary, skills and bullets.
9. NO EMBELLISHMENT. Mirror the advert's TASK words only where the master shows the same task. Never borrow its personality/manner words - do not add calm, compassionate, polite, proactive, strict, fast-paced, under pressure, multidisciplinary, dynamic, passionate, or any similar quality or situation that is not written in the master for that role.

Return ONLY a JSON object, no prose, matching exactly:
{
  "summary": "2-3 sentence role-specific summary drawn only from the master",
  "skills": ["6-10 job-relevant, evidenced skills"],
  "experience": [
    {"title": "...", "organisation": "...", "dates": "...",
     "bullets": ["reordered / lightly reworded existing bullets, max 3"]}
  ],
  "projects": [
    {"name": "...", "bullets": ["..."]}
  ],
  "volunteering": [
    {"title": "...", "organisation": "...", "dates": "...", "bullets": ["..."]}
  ],
  "education": ["one line per qualification, copied essentially verbatim from the master"],
  "certifications": ["only if clearly relevant AND confirmed in the master"],
  "jd_match": [
    {"requirement": "requirement quoted/paraphrased from the advert",
     "essential": true,
     "evidence": "the truthful evidence from the master, or the single word GAP"}
  ],
  "notes_for_human": ["roles omitted and why; anything unsure"]
}
"""


@dataclass
class TailoredResume:
    summary: str = ""
    skills: list[str] = field(default_factory=list)
    experience: list[dict] = field(default_factory=list)
    projects: list[dict] = field(default_factory=list)
    education: list[str] = field(default_factory=list)
    certifications: list[str] = field(default_factory=list)
    volunteering: list[dict] = field(default_factory=list)
    jd_match: list[dict] = field(default_factory=list)
    notes_for_human: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    is_tech: bool = False

    def all_bullets(self) -> list[str]:
        out: list[str] = []
        for entry in self.experience + self.volunteering:
            out.extend(entry.get("bullets", []))
        for entry in self.projects:
            out.extend(entry.get("bullets", []))
        return out

    def fit(self) -> tuple[int, int, list[str]]:
        """(essentials met, essentials total, essential gaps)."""
        ess = [m for m in self.jd_match if m.get("essential")]
        gaps = [m.get("requirement", "") for m in ess
                if str(m.get("evidence", "")).strip().upper().startswith("GAP")]
        return len(ess) - len(gaps), len(ess), gaps

    def gaps(self) -> list[str]:
        return [
            m.get("requirement", "")
            for m in self.jd_match
            if str(m.get("evidence", "")).strip().upper().startswith("GAP")
        ]


def _parse_json(text: str) -> dict:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?|```$", "", text, flags=re.MULTILINE).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end != -1:
            return json.loads(text[start : end + 1])
        raise


def tailor_resume(job: JobPosting, blocks: list[Block] | None = None) -> TailoredResume:
    if not CONFIG.anthropic_api_key:
        raise RuntimeError("ANTHROPIC_API_KEY is not set (see .env / README).")
    blocks = blocks or extract_master()

    master_text = "\n".join(f"[{b.kind}] {b.text}" for b in blocks)
    target_roles = ", ".join(CONFIG.target_roles)
    is_tech = _is_tech_role(job)
    role_type = (
        "TECHNOLOGY / DATA / AI / RESEARCH role - projects, modules and technical skills are relevant."
        if is_tech
        else "NON-TECHNICAL role (customer-facing / admin / care / hospitality / operations) - "
        "lead with people-facing experience; omit projects, programming skills and module lists."
    )

    user = f"""ROLE TYPE: {role_type}

TARGET ROLES (candidate is an early-career MRes AI graduate): {target_roles}

=== CANDIDATE CONSTRAINTS (must obey) ===
{CONFIG.resume_constraints}

=== JOB POSTING ===
Company: {job.company}
Title: {job.title}
URL: {job.url}

Description:
{job.description[:9000]}

=== MASTER RESUME (source of truth - do not go beyond it) ===
{master_text[:22000]}
"""

    client = anthropic.Anthropic(api_key=CONFIG.anthropic_api_key)
    message = client.messages.create(
        model=CONFIG.anthropic_model,
        max_tokens=8000,
        system=TAILOR_SYSTEM,
        messages=[{"role": "user", "content": user}],
    )
    raw = next((b.text for b in message.content if b.type == "text"), "")
    data = _parse_json(raw)

    resume = TailoredResume(
        summary=data.get("summary", ""),
        skills=data.get("skills", []),
        experience=data.get("experience", []),
        projects=data.get("projects", []),
        education=data.get("education", []),
        certifications=data.get("certifications", []),
        volunteering=data.get("volunteering", []),
        jd_match=data.get("jd_match", []),
        notes_for_human=data.get("notes_for_human", []),
        is_tech=is_tech,
    )
    _truth_audit(resume, master_text)
    _apply_constraints(resume)
    _recruiter_first_filters(resume)
    resume.warnings = _fabrication_check(resume, master_corpus(blocks))
    for gap in resume.gaps():
        resume.notes_for_human.append(f"GAP (advert asks, master has no evidence): {gap}")
    return resume




# --------------------------------------------------------------------------- #
# Truth audit: a second, adversarial pass that strips anything the master lacks
# --------------------------------------------------------------------------- #
AUDIT_SYSTEM = """You are a strict CV fact-checker. You receive a candidate's MASTER RESUME (the ONLY source of truth) and a DRAFT CV as JSON. Correct the draft so that every word is supported by the master:
- Delete or reword any claim, adjective, quality, manner or situation that the master does not state for that role - e.g. calm, compassionate, polite, proactive, strict, fast-paced, under pressure, multidisciplinary, dynamic, 'keeping work running smoothly'.
- Keep an advert keyword ONLY where the master supports the same task or fact.
- Never add facts, bullets, roles or skills. You may shorten. Keep titles and dates exactly.
- Keep the same JSON keys and ordering.
Return ONLY the corrected JSON object with one extra key "removed": a short list of what you removed or reworded."""

_AUDIT_KEYS = ("summary", "skills", "experience", "projects", "volunteering", "education", "certifications")


def _truth_audit(resume: TailoredResume, master_text: str) -> None:
    draft = {k: getattr(resume, k) for k in _AUDIT_KEYS}
    client = anthropic.Anthropic(api_key=CONFIG.anthropic_api_key)
    prompt = f"""MASTER RESUME:
{master_text[:22000]}

DRAFT CV JSON:
{json.dumps(draft, ensure_ascii=False)}"""
    try:
        msg = client.messages.create(
            model=CONFIG.anthropic_model,
            max_tokens=6000,
            system=AUDIT_SYSTEM,
            messages=[{"role": "user", "content": prompt}],
        )
        fixed = _parse_json(next((b.text for b in msg.content if b.type == "text"), ""))
    except Exception as exc:  # noqa: BLE001 - never lose the CV because the audit failed
        resume.notes_for_human.append(f"AUDIT SKIPPED (error: {exc}) - read every line before sending")
        return
    for k in _AUDIT_KEYS:
        if k in fixed and fixed[k] not in (None, ""):
            setattr(resume, k, fixed[k])
    for item in fixed.get("removed", [])[:12]:
        resume.notes_for_human.append(f"AUDIT removed/reworded: {item}")


# Manner/quality words that a CV generator tends to borrow from adverts. If one appears in
# the output but nowhere in the master, it is flagged (lookarounds instead of word-boundary
# escapes so the pattern stays plain ASCII).
_TRAIT_RE = re.compile(
    r"(?<![a-z])(calm|compassionate|polite(ly)?|proactive|strict(ly)?|fast[- ]paced|"
    r"multi-?disciplinary|under pressure|dynamic|passionate|enthusiastic|hard[- ]?working|"
    r"energetic|diligent|meticulous|caring|courteous|driven)(?![a-z])",
    re.IGNORECASE,
)



# --------------------------------------------------------------------------- #
# Role type + recruiter-first post-filters (belt-and-braces for the prompt rules)
# --------------------------------------------------------------------------- #
_TECH_TITLE_RE = re.compile(
    r"software|developer|programmer|data (scientist|analyst|engineer)|machine learning|"
    r"\bml\b|\bai\b|artificial intelligence|python|research (assistant|associate|officer)|"
    r"(?<![a-z])it (support|technician|assistant|helpdesk|analyst)|service desk|devops|cyber|"
    r"analytics|engineer",
    re.IGNORECASE,
)
_TECH_SKILL_RE = re.compile(
    r"python|git|github|jupyter|vs ?code|tensorflow|keras|aws|machine learning|"
    r"deep learning|neural|cnn|nlp|programming|pytorch|scikit",
    re.IGNORECASE,
)
_PLACEHOLDER_RE = re.compile(r"\[\s*tbc|tbc\s*\]|\bTBC\b", re.IGNORECASE)
_DEGREE_PITCH_RE = re.compile(r"b\.?\s?agric|distinction|bachelor", re.IGNORECASE)


def _is_tech_role(job: JobPosting) -> bool:
    """A role counts as technical if the TITLE says so (the advert body is too noisy)."""
    return bool(_TECH_TITLE_RE.search(job.title or ""))


def _drop_sentences(text: str, pattern: re.Pattern) -> str:
    parts = re.split(r"(?<=[.!?])\s+", text.strip())
    return " ".join(p for p in parts if not pattern.search(p)).strip()


def _recruiter_first_filters(resume: TailoredResume) -> None:
    def no_placeholder(items):
        return [x for x in items if not _PLACEHOLDER_RE.search(x)]

    # Never leak unconfirmed [TBC] master placeholders onto a CV.
    resume.skills = no_placeholder(resume.skills)
    resume.certifications = no_placeholder(resume.certifications)
    resume.education = no_placeholder(resume.education)
    for entry in resume.experience + resume.projects + resume.volunteering:
        entry["bullets"] = no_placeholder(entry.get("bullets", []))
    if _PLACEHOLDER_RE.search(resume.summary):
        resume.summary = _drop_sentences(resume.summary, _PLACEHOLDER_RE)

    # The undergraduate degree is never a summary pitch.
    resume.summary = _drop_sentences(resume.summary, _DEGREE_PITCH_RE)

    # Roles with no bullets after filtering are dropped rather than printed empty.
    resume.experience = [e for e in resume.experience if e.get("bullets")]
    resume.volunteering = [e for e in resume.volunteering if e.get("bullets")]
    resume.projects = [e for e in resume.projects if e.get("bullets")]

    if not resume.is_tech:
        resume.projects = []
        resume.skills = [s for s in resume.skills if not _TECH_SKILL_RE.search(s)]
        # Drop module lists and dissertation detail from education lines.
        cleaned = []
        for line in resume.education:
            line = re.sub(r"\s*Modules?:.*?(?=(\.\s|$))", "", line, flags=re.IGNORECASE)
            line = re.sub(r"\s*Dissertation in progress.*", "", line, flags=re.IGNORECASE)
            cleaned.append(line.strip(" ."))
        resume.education = cleaned
    resume.skills = resume.skills[:10]


# Belt-and-braces enforcement of the user's corrections, even if the model slips.
#   RESUME_REPLACEMENTS  - "old text=>new text" pairs, ';'-separated (literal)
#   RESUME_DROP_PATTERNS - regex; any skill/bullet/cert line matching it is removed
def _replacements() -> list[tuple[str, str]]:
    out = []
    for pair in CONFIG.resume_replacements.split(";"):
        if "=>" in pair:
            old, new = pair.split("=>", 1)
            out.append((old.strip(), new.strip()))
    return out


def _drop_re():
    pat = CONFIG.resume_drop_patterns.strip()
    return re.compile(pat, re.IGNORECASE) if pat else None


def _rewrite(text: str, pairs: list[tuple[str, str]]) -> str:
    for old, new in pairs:
        if old:
            text = re.sub(re.escape(old), new, text, flags=re.IGNORECASE)
    return text


def _apply_constraints(resume: TailoredResume) -> None:
    pairs = _replacements()
    drop = _drop_re()

    def clean(items: list[str]) -> list[str]:
        kept = [x for x in items if not (drop and drop.search(x))]
        return [_rewrite(x, pairs) for x in kept]

    resume.summary = _rewrite(resume.summary, pairs).strip()
    resume.certifications = clean(resume.certifications)
    resume.skills = clean(resume.skills)
    resume.education = clean(resume.education)
    for entry in resume.experience + resume.projects:
        if "bullets" in entry:
            entry["bullets"] = clean(entry["bullets"])


# --------------------------------------------------------------------------- #
# 3. Guardrail: catch fabricated numbers / bullets with no master source
# --------------------------------------------------------------------------- #
def _numbers(text: str) -> set[str]:
    out = set()
    for m in NUM_RE.finditer(text):
        tok = m.group(0).strip().lower().rstrip(".,").strip()
        if tok and any(ch.isdigit() for ch in tok):
            out.add(tok)
    return out


def _best_ratio(candidate: str, corpus: list[str]) -> float:
    best = 0.0
    for source in corpus:
        r = difflib.SequenceMatcher(None, candidate.lower(), source.lower()).ratio()
        best = max(best, r)
    return best


def _fabrication_check(resume: TailoredResume, corpus: list[str]) -> list[str]:
    warnings: list[str] = []
    source_numbers: set[str] = set()
    for line in corpus:
        source_numbers |= _numbers(line)

    for text in [resume.summary, *resume.skills, *resume.all_bullets()]:
        if not text:
            continue
        new_nums = _numbers(text) - source_numbers
        # allow bare years like 2024 that appear anywhere in the corpus text
        new_nums = {n for n in new_nums if n not in "".join(corpus).lower()}
        if new_nums:
            warnings.append(f"NEW NUMBER not in master: {sorted(new_nums)} -> {text!r}")

    master_lower = " ".join(corpus).lower()
    for text in [resume.summary, *resume.skills, *resume.all_bullets()]:
        for m in _TRAIT_RE.finditer(text or ""):
            if m.group(0).lower() not in master_lower:
                warnings.append(f"UNSUPPORTED WORD {m.group(0)!r} not in master -> {text!r}")

    for bullet in resume.all_bullets():
        ratio = _best_ratio(bullet, corpus)
        if ratio < SIMILARITY_FLOOR:
            warnings.append(
                f"LOW MATCH to master ({ratio:.2f}) - verify this is a real rephrase: {bullet!r}"
            )
    return warnings


# --------------------------------------------------------------------------- #
# 4. Render the tailored resume to a clean, ATS-friendly, ONE-PAGE .docx
# --------------------------------------------------------------------------- #
def _bottom_border(paragraph) -> None:
    """Thin rule under a section heading (renders cleanly in ATS + LibreOffice)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")
    borders.append(bottom)
    p_pr.append(borders)


def _no_dashes(text: str) -> str:
    """David's rule (10 Oct 2026): no em/en dashes in anything he submits.
    En dash marks a range ("2016 – 2021" -> "2016 to 2021"); em dash is a separator (-> comma)."""
    if not isinstance(text, str):
        return text
    text = re.sub(r"\s*–\s*", " to ", text)
    text = re.sub(r"\s*—\s*", ", ", text)
    text = text.replace(" - ", ", ")
    return re.sub(r" {2,}", " ", text)


def _strip_dashes(resume: "TailoredResume") -> None:
    resume.summary = _no_dashes(resume.summary)
    resume.skills = [_no_dashes(x) for x in resume.skills]
    resume.education = [_no_dashes(x) for x in resume.education]
    resume.certifications = [_no_dashes(x) for x in resume.certifications]
    for entry in resume.experience + resume.volunteering + resume.projects:
        for k in ("title", "organisation", "dates", "name"):
            if entry.get(k):
                entry[k] = _no_dashes(entry[k])
        entry["bullets"] = [_no_dashes(b) for b in entry.get("bullets", [])]


def build_docx(resume: TailoredResume, job: JobPosting, out_path, compact: bool = False) -> str:
    """Build the CV. `compact=True` shrinks type/margins slightly to fit one page."""
    from docx.shared import Inches

    _strip_dashes(resume)
    body_pt = 9.5 if compact else 10
    margin = Inches(0.55 if compact else 0.7)

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = margin
        section.left_margin = section.right_margin = margin

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(body_pt)
    style.paragraph_format.space_after = Pt(1)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.0
    bullet_style = doc.styles["List Bullet"]
    bullet_style.font.size = Pt(body_pt)
    bullet_style.paragraph_format.space_after = Pt(0)

    app = CONFIG.applicant

    name = doc.add_paragraph()
    run = name.add_run(app.name)
    run.bold = True
    run.font.size = Pt(17 if not compact else 16)

    contact_bits = [b for b in (app.location, app.email, app.phone) if b]
    if contact_bits:
        doc.add_paragraph(" | ".join(contact_bits)).runs[0].font.size = Pt(9)
    # Links only on technology roles, only when enabled, and never a guessed handle.
    if resume.is_tech and CONFIG.cv_show_links:
        link_bits = [b for b in (app.linkedin, app.github) if b]
        if link_bits:
            doc.add_paragraph(" | ".join(link_bits)).runs[0].font.size = Pt(9)
    # Right-to-work / availability lines are printed ONLY if the candidate configured them.
    status_bits = [b for b in (app.right_to_work, app.availability) if b]
    if status_bits:
        doc.add_paragraph(" | ".join(status_bits)).runs[0].font.size = Pt(9)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(10.5)
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        _bottom_border(p)

    def role_line(entry: dict) -> None:
        head = ", ".join(b for b in (entry.get("title"), entry.get("organisation")) if b)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(head)
        r.bold = True
        if entry.get("dates"):
            p.add_run(f"   |   {entry['dates']}").italic = True

    if resume.summary:
        heading("Profile")
        doc.add_paragraph(resume.summary)

    if resume.skills:
        heading("Key skills")
        doc.add_paragraph(" • ".join(resume.skills))

    if resume.experience:
        heading("Experience")
        for entry in resume.experience:
            role_line(entry)
            for bullet in entry.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    if resume.projects:
        heading("Projects")
        for entry in resume.projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.add_run(entry.get("name", "")).bold = True
            for bullet in entry.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    if resume.volunteering:
        heading("Volunteering")
        for entry in resume.volunteering:
            role_line(entry)
            for bullet in entry.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    if resume.education:
        heading("Education")
        for line in resume.education:
            doc.add_paragraph(line)

    if resume.certifications:
        heading("Certifications")
        for line in resume.certifications:
            doc.add_paragraph(line, style="List Bullet")

    out_path = str(out_path)
    doc.save(out_path)
    return out_path
